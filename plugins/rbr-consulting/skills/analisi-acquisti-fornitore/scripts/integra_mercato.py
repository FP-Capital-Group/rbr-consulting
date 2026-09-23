#!/usr/bin/env python3
"""
Step 4 della skill analisi-acquisti-fornitore (opzionale, dopo costruisci_analisi.py).

Aggiunge i fogli "Confronto Mercato" e "Fonti e Metodo" a un workbook già
costruito, a partire da un JSON con i risultati della ricerca di mercato
(fatta da agenti/ricerche web separate — vedi SKILL.md e
references/prompt_ricerca_mercato.md per lo schema esatto e come strutturare
quella ricerca).

Uso:
    python3 integra_mercato.py Analisi_Acquisti_ROVAGNATI.xlsx mercato_rovagnati.json \
        --mese-rif Giu-26

Il JSON è una lista di oggetti con AL MINIMO: descrizione, prezzo_attuale,
unita ("kg" o "pz"), prezzo_mercato (o null), fonte, tipo_fonte, url,
data_rilevazione, affidabilita, note. Lo script recupera da solo Q.tà/Spesa
del periodo per ciascun articolo leggendoli dal foglio "Articoli" già
presente nel file, facendo un match fuzzy sulla descrizione (le descrizioni
scritte da un agente di ricerca non sono mai identiche al carattere alla
descrizione originale di fattura). Le righe del JSON che NON matchano nessun
articolo del workbook (perché appartengono a un ALTRO fornitore, es. hai
lanciato una ricerca unica per più fornitori piccoli insieme) vengono
scartate con un avviso — non finiscono nel foglio sbagliato.
"""
import argparse
import json
import re
import sys
from difflib import SequenceMatcher

from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

NAVY = "1F3864"
HEADER_FONT = Font(bold=True, color="FFFFFF", size=10)
from openpyxl.styles import PatternFill
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
TITLE_FONT = Font(bold=True, size=16, color=NAVY)
EUR_FMT = '#,##0.00 €'
PCT_FMT = '0.0%'


def norm(s):
    s = s.upper()
    s = re.sub(r'\(.*?\)', '', s)  # rimuove spiegazioni tra parentesi aggiunte dalla ricerca
    s = re.sub(r'[^A-Z0-9]', '', s)
    return s


def find_articolo(desc, articoli, threshold=0.55):
    nd = norm(desc)
    best, best_score = None, 0.0
    for art in articoli:
        npd = norm(art['descrizione'])
        if not npd or not nd:
            continue
        if npd in nd or nd in npd:
            return art
        score = SequenceMatcher(None, nd[:len(npd) + 10], npd).ratio()
        if score > best_score:
            best, best_score = art, score
    return best if best_score >= threshold else None


def valuta(prezzo_attuale, prezzo_mercato):
    if not prezzo_mercato:
        return None
    diff_pct = (prezzo_attuale - prezzo_mercato) / prezzo_mercato
    if diff_pct > 0.08:
        return 'SOPRA MERCATO'
    if diff_pct < -0.08:
        return 'SOTTO MERCATO'
    return 'IN LINEA'


def style_header_row(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)


def autosize(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def read_articoli(path):
    wb = load_workbook(path, data_only=True)
    ws = wb['Articoli']
    headers = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(headers)}
    out = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[idx['Descrizione articolo']] is None:
            continue
        out.append({
            'codice': row[idx['Codice']],
            'descrizione': row[idx['Descrizione articolo']],
            'qta_periodo': row[idx['Q.tà acquistata']],
            'importo_periodo': row[idx['Importo acquisti €']],
        })
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('xlsx')
    ap.add_argument('mercato_json')
    ap.add_argument('--mese-rif', default='ultimo mese disponibile')
    args = ap.parse_args()

    with open(args.mercato_json) as fh:
        mercato_rows = json.load(fh)

    articoli = read_articoli(args.xlsx)

    rows = []
    scartate = []
    for m in mercato_rows:
        art = find_articolo(m['descrizione'], articoli)
        if art is None:
            scartate.append(m['descrizione'])
            continue
        pa = m.get('prezzo_attuale')
        pm = m.get('prezzo_mercato')
        diff_eur = (pa - pm) if (pa is not None and pm) else None
        diff_pct = (diff_eur / pm) if (diff_eur is not None and pm) else None
        importo_periodo = art.get('importo_periodo')
        # Impatto = Diff€ × (Spesa periodo € / Prezzo attuale): robusto anche
        # quando prezzo_attuale è stato normalizzato a un'unità diversa da
        # quella con cui l'articolo è fatturato (es. €/kg da un articolo
        # fatturato a sacchi) — Diff€ × Q.tà fattura si romperebbe in quel caso.
        impatto = (diff_eur * (importo_periodo / pa)) if (diff_eur is not None and pa and importo_periodo) else None
        rows.append({
            'descrizione': art['descrizione'], 'um': m.get('unita', '').lower(), 'prezzo_attuale': pa,
            'prezzo_mercato': pm, 'diff_eur': diff_eur, 'diff_pct': diff_pct, 'impatto': impatto,
            'valutazione': valuta(pa, pm), 'qta_periodo': art.get('qta_periodo'),
            'importo_periodo': importo_periodo, 'fonte': m.get('fonte'), 'tipo_fonte': m.get('tipo_fonte'),
            'url': m.get('url'), 'data_rilevazione': m.get('data_rilevazione'),
            'affidabilita': m.get('affidabilita'), 'note': m.get('note'),
        })

    if scartate:
        print(f"⚠️  {len(scartate)} righe del JSON non corrispondono a nessun articolo di questo fornitore "
              f"(escluse dal foglio — probabilmente appartengono a un altro fornitore della stessa ricerca):")
        for s in scartate:
            print(f"   {s}")

    wb = load_workbook(args.xlsx)
    for name in ('Confronto Mercato', 'Fonti e Metodo'):
        if name in wb.sheetnames:
            del wb[name]

    ws = wb.create_sheet('Confronto Mercato')
    ws.sheet_view.showGridLines = False
    n_con = len([r for r in rows if r['prezzo_mercato']])
    ws.cell(row=1, column=1, value='PREZZI ATTUALI vs PREZZO MEDIO DI MERCATO').font = TITLE_FONT
    ws.cell(row=2, column=1, value=f"Prezzi di {args.mese_rif} (IVA esclusa) confrontati con listini/quotazioni "
                                    f"pubbliche reperite online.").font = Font(size=10)
    ws.cell(row=3, column=1, value=f"Articoli analizzati: {len(rows)} — con benchmark verificato: {n_con} — "
                                    f"senza benchmark pubblico: {len(rows) - n_con} (vedi foglio 'Fonti e Metodo' "
                                    f"e colonna Note).").font = Font(size=9)
    header_row = 5
    cols = ['Descrizione', 'Um', 'Prezzo attuale', 'Prezzo mercato', 'Diff. €', 'Diff. %', 'Impatto sul periodo €',
            'Valutazione', 'Q.tà periodo', 'Spesa periodo €', 'Fonte', 'Tipo fonte', 'Data', 'Affidabilità',
            'Note / limiti del confronto']
    for i, h in enumerate(cols, start=1):
        ws.cell(row=header_row, column=i, value=h)
    style_header_row(ws, header_row, len(cols))
    for i, r in enumerate(rows, start=header_row + 1):
        vals = [r['descrizione'], r['um'], r['prezzo_attuale'], r['prezzo_mercato'], r['diff_eur'], r['diff_pct'],
                r['impatto'], r['valutazione'], r['qta_periodo'], r['importo_periodo'], r['fonte'], r['tipo_fonte'],
                r['data_rilevazione'], r['affidabilita'], r['note']]
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(row=i, column=c, value=v)
            if c in (3, 4, 5, 7, 10) and v is not None:
                cell.number_format = EUR_FMT
            if c == 6 and v is not None:
                cell.number_format = PCT_FMT
            if c == 11 and r['url']:
                cell.hyperlink = r['url']
                cell.font = Font(color='0563C1', underline='single')
        if r['valutazione'] == 'SOPRA MERCATO':
            ws.cell(row=i, column=8).font = Font(bold=True, color='C00000')
        elif r['valutazione'] == 'SOTTO MERCATO':
            ws.cell(row=i, column=8).font = Font(bold=True, color='1F7A1F')
    autosize(ws, [42, 6, 13, 13, 10, 9, 16, 14, 11, 14, 22, 20, 10, 11, 55])
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1).coordinate

    ws2 = wb.create_sheet('Fonti e Metodo')
    ws2.sheet_view.showGridLines = False
    ws2.cell(row=2, column=2, value='FONTI DEI PREZZI DI MERCATO').font = TITLE_FONT
    for i, h in enumerate(['Fonte', 'Tipo', 'Indirizzo', 'Rilevazione']):
        ws2.cell(row=4, column=2 + i, value=h).font = Font(bold=True)
    seen = set()
    row = 5
    for r in rows:
        key = (r['fonte'], r['url'])
        if key in seen or not r['fonte']:
            continue
        seen.add(key)
        ws2.cell(row=row, column=2, value=r['fonte'])
        ws2.cell(row=row, column=3, value=r['tipo_fonte'])
        c = ws2.cell(row=row, column=4, value=r['url'])
        if r['url']:
            c.font = Font(color='0563C1', underline='single')
        ws2.cell(row=row, column=5, value=r['data_rilevazione'])
        row += 1
    autosize(ws2, [4, 34, 26, 45, 12])

    wb.save(args.xlsx)
    print(f"\n{args.xlsx}: aggiunti fogli 'Confronto Mercato' ({len(rows)} articoli, {n_con} con benchmark) "
          f"e 'Fonti e Metodo'.")


if __name__ == '__main__':
    main()
