#!/usr/bin/env python3
"""
Step 2 della skill analisi-acquisti-fornitore.

Costruisce il workbook "Analisi Acquisti Fornitore" (6 fogli: Riepilogo
Fornitore, Articoli, Storico Prezzi, Storico Quantità, Resi - Note di
credito, Dettaglio Righe) a partire dai file fattura di UN fornitore già
selezionati con estrai_fornitori.py.

Uso:
    python3 costruisci_analisi.py <cartella> <file1.xml> [<file2.p7m> ...] \
        --fornitore "ROVAGNATI S.P.A." --piva 00682130968 \
        --cliente "NOME CLIENTE" --out Analisi_Acquisti_ROVAGNATI.xlsx

Più comodo passare i file da una lista (uno per riga) con --lista-file
invece di elencarli tutti sulla riga di comando.

I fogli "Confronto Mercato" e "Fonti e Metodo" (che richiedono ricerca web)
si aggiungono DOPO con integra_mercato.py — questo script produce già un
file .xlsx valido e completo anche senza quei due fogli, se la ricerca di
mercato non serve o va fatta in un secondo momento.
"""
import argparse
import os
import sys
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from fatture_lib import load_invoices, mese_label, is_credito  # noqa: E402

# --- stile RBR (navy/rosso) ---------------------------------------------
NAVY = "1F3864"
RED = "C00000"
GREY = "F2F2F2"
HEADER_FONT = Font(bold=True, color="FFFFFF", size=10)
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
TITLE_FONT = Font(bold=True, size=16, color=NAVY)
SUBTITLE_FONT = Font(bold=True, size=12, color=RED)
SECTION_FONT = Font(bold=True, size=12, color=NAVY)
TOTAL_FILL = PatternFill("solid", fgColor=GREY)
EUR_FMT = '#,##0.00 €'
EUR4_FMT = '#,##0.0000 €'
PCT_FMT = '0.0%'


def style_header_row(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)


def autosize(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


# --- aggregazione ---------------------------------------------------------

def compute_all(docs):
    righe_flat = []
    for d in docs:
        mese = mese_label(d['data'])
        for r in d['righe']:
            righe_flat.append({**r, 'tipo_doc': d['tipo_doc'], 'numero': d['numero'],
                                'data': d['data'], 'mese': mese, 'is_credito': is_credito(d['tipo_doc'])})

    fatture_righe = [r for r in righe_flat if not r['is_credito']]
    credito_righe = [r for r in righe_flat if r['is_credito']]

    mesi_presenti = []
    for d in docs:
        m = mese_label(d['data'])
        if m not in mesi_presenti:
            mesi_presenti.append(m)

    articoli = defaultdict(lambda: {
        'descrizione': None, 'um': None, 'qta_acq': 0.0, 'imp_acq': 0.0,
        'qta_resa': 0.0, 'imp_resa': 0.0, 'prezzi_riga': [], 'mesi': set(), 'n_righe': 0,
        'per_mese_qta': defaultdict(float), 'per_mese_imp': defaultdict(float),
    })

    for r in fatture_righe:
        a = articoli[(r['codice'], r['descrizione'])]
        a['descrizione'] = r['descrizione']
        a['um'] = a['um'] or r['um']
        a['qta_acq'] += r['quantita']
        a['imp_acq'] += r['prezzo_totale']
        a['n_righe'] += 1
        a['mesi'].add(r['mese'])
        if r['quantita']:
            a['prezzi_riga'].append(r['prezzo_totale'] / r['quantita'])
        a['per_mese_qta'][r['mese']] += r['quantita']
        a['per_mese_imp'][r['mese']] += r['prezzo_totale']

    for r in credito_righe:
        a = articoli[(r['codice'], r['descrizione'])]
        a['descrizione'] = a['descrizione'] or r['descrizione']
        a['um'] = a['um'] or r['um']
        a['qta_resa'] += r['quantita']
        a['imp_resa'] += r['prezzo_totale']
        a['n_righe'] += 1

    articoli_rows, storico_prezzi_rows, storico_qta_rows = [], [], []

    for (codice, desc), a in articoli.items():
        qta_netta = a['qta_acq'] - a['qta_resa']
        imp_netto = a['imp_acq'] - a['imp_resa']
        prezzo_medio_netto = (imp_netto / qta_netta) if qta_netta else None
        prezzo_medio_lordo = (a['imp_acq'] / a['qta_acq']) if a['qta_acq'] else None
        prezzo_min = min(a['prezzi_riga']) if a['prezzi_riga'] else None
        prezzo_max = max(a['prezzi_riga']) if a['prezzi_riga'] else None

        mesi_articolo = [m for m in mesi_presenti if m in a['mesi']]
        primo_mese = mesi_articolo[0] if mesi_articolo else None
        ultimo_mese = mesi_articolo[-1] if mesi_articolo else None
        prezzo_primo = (a['per_mese_imp'][primo_mese] / a['per_mese_qta'][primo_mese]) \
            if primo_mese and a['per_mese_qta'][primo_mese] else None
        prezzo_ultimo = (a['per_mese_imp'][ultimo_mese] / a['per_mese_qta'][ultimo_mese]) \
            if ultimo_mese and a['per_mese_qta'][ultimo_mese] else None
        var_eur = (prezzo_ultimo - prezzo_primo) if (prezzo_primo is not None and prezzo_ultimo is not None) else None
        var_pct = (var_eur / prezzo_primo) if (var_eur is not None and prezzo_primo) else None
        incid_resi = (a['imp_resa'] / a['imp_acq']) if a['imp_resa'] and a['imp_acq'] else None

        articoli_rows.append({
            'Codice': codice, 'Descrizione articolo': desc, 'UM': a['um'],
            'Q.tà acquistata': round(a['qta_acq'], 4), 'Importo acquisti €': round(a['imp_acq'], 2),
            'Q.tà resa': round(a['qta_resa'], 4) if a['qta_resa'] else None,
            'Importo resi €': round(a['imp_resa'], 2) if a['imp_resa'] else None,
            'Q.tà netta': round(qta_netta, 4), 'Importo netto €': round(imp_netto, 2),
            'PREZZO MEDIO NETTO': prezzo_medio_netto, 'Prezzo medio lordo': prezzo_medio_lordo,
            'Prezzo min': prezzo_min, 'Prezzo max': prezzo_max,
            '1° mese': primo_mese, 'Prezzo 1° mese': prezzo_primo,
            'Ultimo mese': ultimo_mese, 'Prezzo ultimo mese': prezzo_ultimo,
            'Variazione €': var_eur, 'Variazione %': var_pct,
            'Incidenza resi %': incid_resi, 'N. righe': a['n_righe'], 'N. mesi': len(mesi_articolo),
        })

        prezzi_mensili = {m: (a['per_mese_imp'][m] / a['per_mese_qta'][m] if a['per_mese_qta'][m] else None)
                          for m in mesi_presenti}
        row_pz = {'Codice': codice, 'Descrizione articolo': desc, 'UM': a['um'], **prezzi_mensili,
                  'Var. 1°→ultimo %': var_pct}
        storico_prezzi_rows.append(row_pz)

        row_q = {'Codice': codice, 'Descrizione articolo': desc, 'UM': a['um']}
        for m in mesi_presenti:
            row_q[m] = a['per_mese_qta'].get(m) or None
        row_q['Totale q.tà'] = round(a['qta_acq'], 4)
        row_q['Totale €'] = round(a['imp_acq'], 2)
        storico_qta_rows.append(row_q)

    resi_rows = []
    tot_qta_resa = tot_imp_resa = 0.0
    for d in docs:
        if not is_credito(d['tipo_doc']):
            continue
        for r in d['righe']:
            a = articoli.get((r['codice'], r['descrizione']), {})
            qta_acq_periodo = a.get('qta_acq', 0.0)
            imp_acq_periodo = a.get('imp_acq', 0.0)
            resi_rows.append({
                'Nota di credito': d['numero'], 'Data': d['data'], 'Codice': r['codice'],
                'Descrizione articolo': r['descrizione'], 'UM': r['um'], 'Q.tà resa': r['quantita'],
                'Prezzo unit.': (r['prezzo_totale'] / r['quantita']) if r['quantita'] else None,
                'Importo reso €': r['prezzo_totale'],
                'Q.tà acquistata periodo': qta_acq_periodo if qta_acq_periodo else None,
                'Importo acquistato €': imp_acq_periodo if imp_acq_periodo else None,
                'Incidenza reso su acquisti %': (r['prezzo_totale'] / imp_acq_periodo) if imp_acq_periodo else None,
            })
            tot_qta_resa += r['quantita']
            tot_imp_resa += r['prezzo_totale']

    tot_acq = sum(a['imp_acq'] for a in articoli.values())
    if resi_rows:
        resi_rows.append({
            'Nota di credito': 'TOTALE', 'Data': None, 'Codice': None, 'Descrizione articolo': None, 'UM': None,
            'Q.tà resa': round(tot_qta_resa, 4), 'Prezzo unit.': None, 'Importo reso €': round(tot_imp_resa, 2),
            'Q.tà acquistata periodo': None, 'Importo acquistato €': round(tot_acq, 2),
            'Incidenza reso su acquisti %': (tot_imp_resa / tot_acq) if tot_acq else None,
        })

    dettaglio_rows = []
    for d in docs:
        m = mese_label(d['data'])
        for r in d['righe']:
            dettaglio_rows.append({
                'Tipo doc': 'Nota di credito' if is_credito(d['tipo_doc']) else 'Fattura',
                'N. doc': d['numero'], 'Data': d['data'], 'Mese': m,
                'Codice': r['codice'], 'Descrizione articolo': r['descrizione'], 'UM': r['um'],
                'Quantità': r['quantita'],
                'Prezzo unitario': (r['prezzo_totale'] / r['quantita']) if r['quantita'] else None,
                'Importo €': r['prezzo_totale'],
                'IVA': f"{r['iva']:g}%" if r['iva'] is not None else None,
            })

    andamento = []
    for m in mesi_presenti:
        acq = sum(r['prezzo_totale'] for r in fatture_righe if r['mese'] == m)
        cred = -sum(r['prezzo_totale'] for r in credito_righe if r['mese'] == m)
        n_art = len(set((r['codice'], r['descrizione']) for r in fatture_righe if r['mese'] == m))
        n_righe_m = len([r for r in righe_flat if r['mese'] == m])
        andamento.append({'Mese': m, 'Acquisti €': round(acq, 2), 'Note credito €': round(cred, 2),
                           'Netto €': round(acq - cred, 2), 'Incid. resi %': (cred / acq) if acq else 0,
                           'N. articoli': n_art, 'N. righe': n_righe_m})

    scostamenti = [d for d in docs if abs(d.get('scostamento', 0)) > 0.05]

    return {
        'articoli_rows': articoli_rows, 'storico_prezzi_rows': storico_prezzi_rows,
        'storico_qta_rows': storico_qta_rows, 'resi_rows': resi_rows, 'dettaglio_rows': dettaglio_rows,
        'andamento': andamento, 'mesi_presenti': mesi_presenti, 'tot_acq': tot_acq, 'tot_resi': tot_imp_resa,
        'n_fatture': len([d for d in docs if not is_credito(d['tipo_doc'])]),
        'n_note': len([d for d in docs if is_credito(d['tipo_doc'])]),
        'n_righe_tot': len(righe_flat), 'n_articoli': len(articoli), 'scostamenti': scostamenti,
    }


# --- scrittura fogli --------------------------------------------------------

def write_riepilogo(wb, computed, supplier_name, supplier_piva, client_name):
    ws = wb.active
    ws.title = 'Riepilogo Fornitore'
    ws.sheet_view.showGridLines = False
    autosize(ws, [4, 42, 18, 14, 14, 14, 14, 14])

    mesi = computed['mesi_presenti']
    periodo = f"{mesi[0]} / {mesi[-1]}" if mesi else "-"
    row = 2
    ws.cell(row=row, column=2, value='ANALISI ACQUISTI FORNITORE').font = TITLE_FONT
    row += 1
    ws.cell(row=row, column=2, value=f"{supplier_name} — P.IVA {supplier_piva}").font = SUBTITLE_FONT
    row += 1
    ws.cell(row=row, column=2, value=f"Cliente: {client_name} — Periodo: {periodo}")
    row += 1
    ws.cell(row=row, column=2,
            value=f"Documenti analizzati: {computed['n_fatture']} fatture + {computed['n_note']} note di credito — "
                  f"{computed['n_righe_tot']} righe — {computed['n_articoli']} articoli distinti").font = Font(size=9)
    row += 1
    if computed['scostamenti']:
        tot_scost = sum(d['scostamento'] for d in computed['scostamenti'])
        ws.cell(row=row, column=2,
                value=f"⚠️ {len(computed['scostamenti'])} documenti hanno un imponibile dichiarato diverso dalla "
                      f"somma delle righe prodotto (totale scostamento {tot_scost:+.2f}€, spesso sconti finanziari "
                      f"di fattura non legati a un prodotto) — vedi dettaglio in fondo a questo foglio."
                ).font = Font(size=9, italic=True, color=RED)
        row += 1
    row += 1

    ws.cell(row=row, column=2, value='TOTALI DI PERIODO (imponibile, IVA esclusa)').font = SECTION_FONT
    row += 1
    ws.cell(row=row, column=2, value='Totale acquisti (fatture)')
    ws.cell(row=row, column=3, value=round(computed['tot_acq'], 2)).number_format = EUR_FMT
    ws.cell(row=row, column=3).font = Font(bold=True)
    row += 1
    ws.cell(row=row, column=2, value='Totale note di credito (resi/accrediti)')
    ws.cell(row=row, column=3, value=round(-computed['tot_resi'], 2)).number_format = EUR_FMT
    ws.cell(row=row, column=3).font = Font(bold=True)
    row += 1
    netto = computed['tot_acq'] - computed['tot_resi']
    ws.cell(row=row, column=2, value='TOTALE NETTO FORNITORE').font = Font(bold=True)
    ws.cell(row=row, column=2).fill = TOTAL_FILL
    c3 = ws.cell(row=row, column=3, value=round(netto, 2))
    c3.number_format = EUR_FMT
    c3.font = Font(bold=True)
    c3.fill = TOTAL_FILL
    row += 1
    ws.cell(row=row, column=2, value='Incidenza resi sul fatturato')
    c = ws.cell(row=row, column=3, value=(computed['tot_resi'] / computed['tot_acq']) if computed['tot_acq'] else 0)
    c.number_format = PCT_FMT
    c.font = Font(bold=True)
    row += 2

    ws.cell(row=row, column=2, value='ANDAMENTO MENSILE').font = SECTION_FONT
    row += 1
    for i, h in enumerate(['Mese', 'Acquisti €', 'Note credito €', 'Netto €', 'Incid. resi %', 'N. articoli', 'N. righe']):
        ws.cell(row=row, column=2 + i, value=h).font = Font(bold=True)
    row += 1
    for a in computed['andamento']:
        ws.cell(row=row, column=2, value=a['Mese'])
        ws.cell(row=row, column=3, value=a['Acquisti €']).number_format = EUR_FMT
        ws.cell(row=row, column=4, value=a['Note credito €']).number_format = EUR_FMT
        ws.cell(row=row, column=5, value=a['Netto €']).number_format = EUR_FMT
        ws.cell(row=row, column=6, value=a['Incid. resi %']).number_format = PCT_FMT
        ws.cell(row=row, column=7, value=a['N. articoli'])
        ws.cell(row=row, column=8, value=a['N. righe'])
        row += 1
    tot_a = sum(a['Acquisti €'] for a in computed['andamento'])
    tot_c = sum(a['Note credito €'] for a in computed['andamento'])
    ws.cell(row=row, column=2, value='TOTALE').font = Font(bold=True)
    ws.cell(row=row, column=3, value=round(tot_a, 2)).number_format = EUR_FMT
    ws.cell(row=row, column=4, value=round(tot_c, 2)).number_format = EUR_FMT
    ws.cell(row=row, column=5, value=round(tot_a - tot_c, 2)).number_format = EUR_FMT
    ws.cell(row=row, column=6, value=(tot_c / tot_a) if tot_a else 0).number_format = PCT_FMT
    ws.cell(row=row, column=7, value=computed['n_articoli'])
    ws.cell(row=row, column=8, value=computed['n_righe_tot'])
    row += 2

    arts = [a for a in computed['articoli_rows'] if a['Variazione %'] is not None]
    top_up = sorted(arts, key=lambda a: -a['Variazione %'])[:10]
    top_down = sorted(arts, key=lambda a: a['Variazione %'])[:10]
    top_resi = sorted([a for a in computed['articoli_rows'] if a['Importo resi €']],
                       key=lambda a: -a['Importo resi €'])[:10]

    def blocco_top(titolo, righe, colonne, extractor):
        nonlocal row
        ws.cell(row=row, column=2, value=titolo).font = SECTION_FONT
        row += 1
        for i, h in enumerate(colonne):
            ws.cell(row=row, column=2 + i, value=h).font = Font(bold=True)
        row += 1
        for a in righe:
            vals = extractor(a)
            for i, v in enumerate(vals):
                cell = ws.cell(row=row, column=2 + i, value=v)
                if isinstance(v, float) and colonne[i].endswith('%'):
                    cell.number_format = PCT_FMT
                elif isinstance(v, float):
                    cell.number_format = EUR4_FMT
            row += 1
        row += 1

    blocco_top('TOP 10 RINCARI (prezzo primo mese → ultimo mese)', top_up,
                ['Codice', 'Descrizione', 'Prezzo iniziale', 'Prezzo finale', 'Var. %', 'Acquistato €'],
                lambda a: [a['Codice'], a['Descrizione articolo'], a['Prezzo 1° mese'], a['Prezzo ultimo mese'],
                           a['Variazione %'], a['Importo acquisti €']])
    blocco_top('TOP 10 RIBASSI (prezzo primo mese → ultimo mese)', top_down,
                ['Codice', 'Descrizione', 'Prezzo iniziale', 'Prezzo finale', 'Var. %', 'Acquistato €'],
                lambda a: [a['Codice'], a['Descrizione articolo'], a['Prezzo 1° mese'], a['Prezzo ultimo mese'],
                           a['Variazione %'], a['Importo acquisti €']])
    if top_resi:
        blocco_top('TOP 10 ARTICOLI PER RESI', top_resi,
                    ['Codice', 'Descrizione', 'Q.tà resa', 'Reso €', 'Incid. su acquisti %', 'Acquistato €'],
                    lambda a: [a['Codice'], a['Descrizione articolo'], a['Q.tà resa'], a['Importo resi €'],
                               a['Incidenza resi %'], a['Importo acquisti €']])

    if computed['scostamenti']:
        ws.cell(row=row, column=2, value='DOCUMENTI CON SCOSTAMENTO IMPONIBILE DICHIARATO vs RIGHE').font = SECTION_FONT
        row += 1
        for i, h in enumerate(['N. doc', 'Data', 'Imponibile dichiarato €', 'Somma righe €', 'Scostamento €']):
            ws.cell(row=row, column=2 + i, value=h).font = Font(bold=True)
        row += 1
        for d in computed['scostamenti']:
            ws.cell(row=row, column=2, value=d['numero'])
            ws.cell(row=row, column=3, value=d['data'])
            ws.cell(row=row, column=4, value=d['imponibile_dichiarato']).number_format = EUR_FMT
            ws.cell(row=row, column=5, value=d['imponibile_da_righe']).number_format = EUR_FMT
            ws.cell(row=row, column=6, value=d['scostamento']).number_format = EUR_FMT
            row += 1


def write_articoli(wb, computed):
    ws = wb.create_sheet('Articoli')
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = 'A2'
    cols = ['Codice', 'Descrizione articolo', 'UM', 'Q.tà acquistata', 'Importo acquisti €', 'Q.tà resa',
            'Importo resi €', 'Q.tà netta', 'Importo netto €', 'PREZZO MEDIO NETTO', 'Prezzo medio lordo',
            'Prezzo min', 'Prezzo max', '1° mese', 'Prezzo 1° mese', 'Ultimo mese', 'Prezzo ultimo mese',
            'Variazione €', 'Variazione %', 'Incidenza resi %', 'N. righe', 'N. mesi']
    for i, h in enumerate(cols, start=1):
        ws.cell(row=1, column=i, value=h)
    style_header_row(ws, 1, len(cols))
    money_cols = {5, 7, 9, 10, 11, 12, 13, 15, 17, 18}
    pct_cols = {19, 20}
    for r, a in enumerate(sorted(computed['articoli_rows'], key=lambda a: -a['Importo acquisti €']), start=2):
        for i, key in enumerate(cols, start=1):
            v = a[key]
            cell = ws.cell(row=r, column=i, value=v)
            if i in money_cols and v is not None:
                cell.number_format = EUR4_FMT if key in (
                    'PREZZO MEDIO NETTO', 'Prezzo medio lordo', 'Prezzo min', 'Prezzo max', 'Prezzo 1° mese',
                    'Prezzo ultimo mese', 'Variazione €') else EUR_FMT
            if i in pct_cols and v is not None:
                cell.number_format = PCT_FMT
    autosize(ws, [12, 46, 6, 13, 15, 11, 13, 12, 15, 17, 15, 12, 12, 9, 14, 10, 15, 13, 12, 14, 9, 9])


def write_storico_prezzi(wb, computed):
    ws = wb.create_sheet('Storico Prezzi')
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = 'A2'
    cols = ['Codice', 'Descrizione articolo', 'UM'] + computed['mesi_presenti'] + ['Var. 1°→ultimo %']
    for i, h in enumerate(cols, start=1):
        ws.cell(row=1, column=i, value=h)
    style_header_row(ws, 1, len(cols))
    n_mesi = len(computed['mesi_presenti'])
    for r, a in enumerate(sorted(computed['storico_prezzi_rows'], key=lambda a: a['Descrizione articolo'] or ''), start=2):
        for i, key in enumerate(cols, start=1):
            v = a.get(key)
            cell = ws.cell(row=r, column=i, value=v)
            if 4 <= i <= 3 + n_mesi and v is not None:
                cell.number_format = EUR4_FMT
            if key == 'Var. 1°→ultimo %' and v is not None:
                cell.number_format = PCT_FMT
    autosize(ws, [12, 46, 6] + [11] * n_mesi + [16])


def write_storico_qta(wb, computed):
    ws = wb.create_sheet('Storico Quantità')
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = 'A2'
    cols = ['Codice', 'Descrizione articolo', 'UM'] + computed['mesi_presenti'] + ['Totale q.tà', 'Totale €']
    for i, h in enumerate(cols, start=1):
        ws.cell(row=1, column=i, value=h)
    style_header_row(ws, 1, len(cols))
    for r, a in enumerate(sorted(computed['storico_qta_rows'], key=lambda a: a['Descrizione articolo'] or ''), start=2):
        for i, key in enumerate(cols, start=1):
            v = a.get(key)
            cell = ws.cell(row=r, column=i, value=v)
            if key == 'Totale €' and v is not None:
                cell.number_format = EUR_FMT
    n_mesi = len(computed['mesi_presenti'])
    autosize(ws, [12, 46, 6] + [11] * n_mesi + [12, 13])


def write_resi(wb, computed):
    ws = wb.create_sheet('Resi - Note di credito')
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = 'A2'
    cols = ['Nota di credito', 'Data', 'Codice', 'Descrizione articolo', 'UM', 'Q.tà resa', 'Prezzo unit.',
            'Importo reso €', 'Q.tà acquistata periodo', 'Importo acquistato €', 'Incidenza reso su acquisti %']
    for i, h in enumerate(cols, start=1):
        ws.cell(row=1, column=i, value=h)
    style_header_row(ws, 1, len(cols))
    for r, a in enumerate(computed['resi_rows'], start=2):
        for i, key in enumerate(cols, start=1):
            v = a.get(key)
            cell = ws.cell(row=r, column=i, value=v)
            if key in ('Prezzo unit.', 'Importo reso €', 'Importo acquistato €') and v is not None:
                cell.number_format = EUR_FMT
            if key == 'Incidenza reso su acquisti %' and v is not None:
                cell.number_format = PCT_FMT
        if a.get('Nota di credito') == 'TOTALE':
            for i in range(1, len(cols) + 1):
                ws.cell(row=r, column=i).font = Font(bold=True)
                ws.cell(row=r, column=i).fill = TOTAL_FILL
    autosize(ws, [14, 12, 12, 46, 6, 10, 12, 14, 18, 16, 20])


def write_dettaglio(wb, computed):
    ws = wb.create_sheet('Dettaglio Righe')
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = 'A2'
    cols = ['Tipo doc', 'N. doc', 'Data', 'Mese', 'Codice', 'Descrizione articolo', 'UM', 'Quantità',
            'Prezzo unitario', 'Importo €', 'IVA']
    for i, h in enumerate(cols, start=1):
        ws.cell(row=1, column=i, value=h)
    style_header_row(ws, 1, len(cols))
    for r, a in enumerate(computed['dettaglio_rows'], start=2):
        for i, key in enumerate(cols, start=1):
            v = a.get(key)
            cell = ws.cell(row=r, column=i, value=v)
            if key == 'Prezzo unitario' and v is not None:
                cell.number_format = EUR4_FMT
            if key == 'Importo €' and v is not None:
                cell.number_format = EUR_FMT
    autosize(ws, [16, 14, 12, 8, 12, 46, 6, 10, 13, 12, 6])


def build_workbook(docs, supplier_name, supplier_piva, client_name, out_path):
    computed = compute_all(docs)
    wb = Workbook()
    write_riepilogo(wb, computed, supplier_name, supplier_piva, client_name)
    write_articoli(wb, computed)
    write_storico_prezzi(wb, computed)
    write_storico_qta(wb, computed)
    write_resi(wb, computed)
    write_dettaglio(wb, computed)
    wb.save(out_path)
    return computed


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('cartella')
    ap.add_argument('files', nargs='*', help='Nomi file (dentro <cartella>) da includere')
    ap.add_argument('--lista-file', help='Path a un .txt con un nome file per riga, alternativa a elencarli')
    ap.add_argument('--fornitore', required=True)
    ap.add_argument('--piva', required=True)
    ap.add_argument('--cliente', required=True)
    ap.add_argument('--out', required=True)
    args = ap.parse_args()

    files = list(args.files)
    if args.lista_file:
        with open(args.lista_file) as fh:
            files += [line.strip() for line in fh if line.strip()]
    if not files:
        print("Nessun file specificato (usa gli argomenti posizionali o --lista-file)", file=sys.stderr)
        sys.exit(1)

    paths = [os.path.join(args.cartella, f) for f in files]
    docs, errori = load_invoices(paths)
    if errori:
        print(f"⚠️  {len(errori)} file non leggibili, esclusi dall'analisi:")
        for e in errori:
            print(f"   {e}")

    computed = build_workbook(docs, args.fornitore, args.piva, args.cliente, args.out)
    print(f"\n{args.out} creato: {computed['n_articoli']} articoli, {computed['n_righe_tot']} righe, "
          f"totale acquisti {computed['tot_acq']:.2f}€, totale resi {computed['tot_resi']:.2f}€")
    if computed['scostamenti']:
        print(f"⚠️  {len(computed['scostamenti'])} documenti con scostamento imponibile dichiarato vs righe "
              f"(dettaglio nel foglio Riepilogo Fornitore)")


if __name__ == '__main__':
    main()
